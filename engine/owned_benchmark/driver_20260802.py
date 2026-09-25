#!/usr/bin/env python3
"""Driver for 2026-08-02 scheduled fire. Implements task sequencing rules:
providers strictly one at a time (anthropic, deepseek, mistral, google, openai);
stop conditions per provider: ALL_DONE / stopped_on_budget / STALL (pending not
falling across 3 consecutive calls) / 70-call ceiling. Zero-throughput validity
re-measure, workbook completion check, status file for polling."""
import json, os, re, subprocess, sys, datetime

DIR = os.path.dirname(os.path.abspath(__file__))
DESKTOP = os.path.dirname(os.path.dirname(DIR))
os.chdir(DIR)
TODAY = datetime.date.today()
STATUS = os.path.join(DIR, "results", "driver_status_20260802.json")

ORDER = [("anthropic", 130), ("deepseek", 90), ("mistral", 130), ("google", 200), ("openai", 300)]
WB = {
    "openai": "InferenceCarbon_TokenSec_Daily_Owned.xlsx",
    "anthropic": "InferenceCarbon_TokenSec_Daily_Anthropic_Owned.xlsx",
    "google": "InferenceCarbon_TokenSec_Daily_Gemini_Owned.xlsx",
    "mistral": "InferenceCarbon_TokenSec_Daily_Mistral_Owned.xlsx",
    "deepseek": "InferenceCarbon_TokenSec_Daily_DeepSeek_Owned.xlsx",
}

state = {"started": datetime.datetime.now().isoformat(), "providers": {}, "current": None, "finished": False}

def save():
    tmp = STATUS + ".tmp"
    with open(tmp, "w") as f:
        json.dump(state, f, indent=1, default=str)
    os.replace(tmp, STATUS)

def last_json(text):
    # find last parseable JSON object in stdout
    idx = [m.start() for m in re.finditer(r"^\{", text, re.M)]
    for i in reversed(idx):
        try:
            return json.loads(text[i:])
        except Exception:
            continue
    return None

def bench(args, timeout=300):
    try:
        r = subprocess.run([sys.executable, "inferencecarbon_bench.py"] + args,
                           capture_output=True, text=True, timeout=timeout)
        return r.stdout + "\n" + r.stderr
    except subprocess.TimeoutExpired as e:
        return (e.stdout or "") + "\nDRIVER_SUBPROCESS_TIMEOUT"

def today_row_check(prov):
    """Return (filled, total, gaps[list of headers blank/zero]) for today's row."""
    try:
        import openpyxl
        path = os.path.join(DESKTOP, WB[prov])
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        headers = [c.value for c in ws[1]][1:]
        row = None
        for r in ws.iter_rows(min_row=2):
            v = r[0].value
            d = v.date() if isinstance(v, datetime.datetime) else v if isinstance(v, datetime.date) else None
            if d is None and v is not None:
                try:
                    d = datetime.datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
                except Exception:
                    d = None
            if d == TODAY:
                row = r
        if row is None:
            return 0, len([h for h in headers if h]), ["NO_TODAY_ROW"]
        gaps = []
        total = 0
        filled = 0
        for h, c in zip(headers, row[1:]):
            if h is None or str(h).strip() == "":
                continue
            total += 1
            v = c.value
            if isinstance(v, (int, float)) and v != 0:
                filled += 1
            else:
                gaps.append(str(h))
        return filled, total, gaps
    except Exception as e:
        return -1, -1, ["CHECK_ERROR: %s" % e]

for prov, cap in ORDER:
    ps = {"cap": cap, "calls": 0, "measured_this_fire": 0, "stop": None, "final": None,
          "new_candidates": None, "zero_remeasured": [], "gaps": None, "row": None}
    state["providers"][prov] = ps
    state["current"] = prov
    save()

    out = bench(["discover", "--provider", prov], timeout=120)
    dj = last_json(out)
    ps["new_candidates"] = (dj or {}).get("new_candidates")
    save()

    pend_hist = []
    while ps["calls"] < 70:
        ps["calls"] += 1
        out = bench(["run", "--provider", prov, "--workload", "10k", "--repeats", "1",
                     "--max-output", "2500", "--budget", str(cap), "--time-budget", "3",
                     "--request-timeout", "40"])
        j = last_json(out)
        if j is None:
            # interrupted/errored call is NOT a stop condition; re-issue
            save()
            continue
        ps["final"] = j
        ps["measured_this_fire"] += j.get("models_this_call", 0) or 0
        pend = j.get("pending")
        if pend is not None:
            pend_hist.append(pend)
        save()
        if j.get("ALL_DONE"):
            ps["stop"] = "ALL_DONE"
            break
        if j.get("stopped_on_budget"):
            ps["stop"] = "BUDGET"
            break
        if len(pend_hist) >= 3 and pend_hist[-1] >= pend_hist[-2] >= pend_hist[-3] and pend_hist[-1] == pend_hist[-3]:
            ps["stop"] = "STALL"
            break
    if ps["stop"] is None:
        ps["stop"] = "CALL_CEILING"
    save()

    # STEP 2c validity: re-measure zeros once (unless stopped on budget)
    filled, total, gaps = today_row_check(prov)
    zero_like = [g for g in gaps if not g.startswith(("NO_TODAY_ROW", "CHECK_ERROR"))]
    if zero_like and ps["stop"] != "BUDGET":
        for g in zero_like[:10]:
            base = re.sub(r"\s*\(.*\)$", "", g).strip()
            ps["zero_remeasured"].append(g)
            out = bench(["run", "--provider", prov, "--workload", "10k", "--repeats", "1",
                         "--max-output", "2500", "--budget", str(cap), "--time-budget", "3",
                         "--request-timeout", "40", "--models", base])
            j = last_json(out)
            if j:
                ps["measured_this_fire"] += j.get("models_this_call", 0) or 0
                ps["final"] = j if j.get("est_cumulative_usd") else ps["final"]
        filled, total, gaps = today_row_check(prov)
    ps["row"] = "%s of %s" % (filled, total)
    ps["gaps"] = gaps
    save()

state["current"] = None
state["finished"] = True
state["ended"] = datetime.datetime.now().isoformat()
save()
print("DRIVER DONE")
