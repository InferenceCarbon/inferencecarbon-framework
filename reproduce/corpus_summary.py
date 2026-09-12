#!/usr/bin/env python3
"""corpus_summary.py — regenerate the workbook's corpus-derived inputs from the frozen corpus.

InferenceCarbon framework v1.0.0. Reproduces, from InferenceCarbon_corpus_frozen_20260822.zip
(SHA-256 01c57eb5bf41476fee8e06a8b677a72c172c849741cc65a9b554cf7397d84cc6):

  corpus.csv      the Corpus sheet of OpenAIModelCalculations (TPS P50, n, daily-median mean/SD,
                  days, R central / short / long / floor / heavy with n)  — paper Appendix B, D.5.5
  table_b1.csv    Table B1 (daily-median TPS mean, SD, days)
  bootstrap.csv   the Bootstrap sheet (T and R bands, B = 10,000)          — paper Appendix D.6
  ratios.csv      the RatioEnvelopes bootstrap (5th/95th of within-family ratios) — Table C3

Filters are the 'cleared-down' set of refresh_tables.py (paper Appendix B.4), restated here so the
deposit is self-contained:
  F1  mock records excluded; provider == openai only.
  F2  reasoning (R) cells: prompt_class in {reason, reason-heavy}, prompt_version == 2 only.
  F3  R cells require ok, reasoning_tokens present, visible (o200k) > 50 and, for the fixed-task
      'reason' cells, min_answer_met (an answer that stopped short of the prompt's word minimum is a
      truncation, not a reasoning measurement; added 2026-09-10). The k = 36 'reason-heavy' rung is
      not subject to the word-minimum test (its answers are puzzle solutions, not summaries) — the
      deposited workbook's heavy column was computed on that basis. Records clamped to
      max_tokens < 2000 are excluded (truncation artefact).
  F4  TPS series: workload 10k, prompt_class summarise, ok, output_tps_o200k present.
  F5  R floors: hidden-token ratio on the 10k summarise records (visible > 50).
  F6  variants with < 8 TPS records excluded from the tables and listed in the report.

Usage:  python corpus_summary.py --corpus InferenceCarbon_corpus_frozen_20260822.zip --out out/
        add --workbook OpenAIModelCalculations_v1.0.0.xlsx to diff against the deposited workbook.
Deterministic: seeds 20260822 (T and R bands) and 20260823 (ratio envelopes); numpy default_rng.
"""
import argparse, csv, glob, json, os, statistics as st, sys, tempfile, zipfile, hashlib
from collections import defaultdict
import numpy as np

CORPUS_SHA = '01c57eb5bf41476fee8e06a8b677a72c172c849741cc65a9b554cf7397d84cc6'
B = 10000
EFF = ['minimal', 'low', 'medium', 'high', 'xhigh']
NONREASON = {'gpt-4o-mini', 'gpt-5 (minimal)', 'gpt-5.2 (minimal)', 'gpt-5.4 (minimal)',
             'gpt-5.4-nano (minimal)', 'gpt-5.4-mini (minimal)'}   # measure R = 1.0 (paper 3.3)
ROW_LABEL = {'gpt-5 (minimal)': 'gpt-5 (non-reasoning)', 'gpt-5.2 (minimal)': 'gpt-5.2 (non-reasoning)',
             'gpt-5.4 (minimal)': 'gpt-5.4 (non-reasoning)', 'gpt-5.4-nano (minimal)': 'gpt-5.4-nano (non-reasoning)',
             'gpt-5.4-mini (minimal)': 'gpt-5.4-mini (non-reasoning)'}
RATIOS = [  # Table C3 comparisons (numerator, denominator)
 ('gpt-5.6-sol (minimal)', 'gpt-5.6-luna (minimal)'), ('gpt-5.6-sol (medium)', 'gpt-5.6-luna (medium)'),
 ('gpt-5.6-sol (xhigh)', 'gpt-5.6-luna (xhigh)'), ('gpt-5.6-sol (xhigh)', 'gpt-5.6-sol (minimal)'),
 ('gpt-5.6-luna (xhigh)', 'gpt-5.6-luna (minimal)'), ('gpt-5.6-terra (xhigh)', 'gpt-5.6-terra (minimal)'),
 ('gpt-5.5 (xhigh)', 'gpt-5.5 (minimal)'), ('gpt-5.4 (xhigh)', 'gpt-5.4 (minimal)'),
 ('gpt-5.4 (xhigh)', 'gpt-5.4-nano (xhigh)'), ('gpt-5.4 (xhigh)', 'gpt-5.4-mini (xhigh)'),
 ('gpt-5.5 (xhigh)', 'gpt-5.6-terra (xhigh)')]

def load(corpus):
    if corpus.endswith('.zip'):
        h = hashlib.sha256(open(corpus, 'rb').read()).hexdigest()
        if h != CORPUS_SHA:
            sys.exit(f'corpus checksum mismatch: {h}')
        tmp = tempfile.mkdtemp(); zipfile.ZipFile(corpus).extractall(tmp); root = tmp
    else:
        root = corpus
    files = sorted(glob.glob(os.path.join(root, '**', 'run_2*.json'), recursive=True))
    tps = defaultdict(list); tps_day = defaultdict(lambda: defaultdict(list))
    R = defaultdict(list); floor = defaultdict(list); excl = defaultdict(int)
    for fp in files:
        try: d = json.load(open(fp))
        except Exception: continue
        if d.get('mock') or d.get('provider', 'openai') != 'openai': continue            # F1
        wl = d.get('workload'); day = (d.get('generated') or d.get('run_id'))[:10].replace('T', '')
        for r in d.get('records', []):
            lab = r.get('label'); pc = r.get('prompt_class', 'summarise')
            v, h = r.get('output_tokens_o200k'), r.get('reasoning_tokens')
            if wl == '10k' and pc == 'summarise' and r.get('ok') and r.get('output_tps_o200k'):
                tps[lab].append(r['output_tps_o200k']); tps_day[lab][day].append(r['output_tps_o200k'])  # F4
                if v and h is not None and v > 50: floor[lab].append((v + h) / v)                      # F5
            if pc in ('reason', 'reason-heavy'):
                cl = r.get('max_tokens_clamped')
                if cl is not None and cl < 2000: excl['truncation_clamp'] += 1; continue
                if r.get('prompt_version') != 2: excl['v1_reason'] += 1; continue                       # F2
                if pc == 'reason' and not r.get('min_answer_met'): excl['min_answer_not_met'] += 1; continue
                if r.get('ok') and v and h is not None and v > 50: R[(lab, wl, pc)].append((v + h) / v)  # F3
                else: excl['censored_or_failed'] += 1
    return tps, tps_day, R, floor, dict(excl)

med = lambda xs: st.median(xs) if xs else None
r2 = lambda x: None if x is None else round(x, 2)

def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--corpus', required=True); ap.add_argument('--out', default='out')
    ap.add_argument('--workbook'); a = ap.parse_args(); os.makedirs(a.out, exist_ok=True)
    tps, tps_day, R, floor, excl = load(a.corpus)
    labels = [l for l in tps if l == 'gpt-4o-mini' or l.startswith('gpt-5')]
    fams = sorted({l.split(' (')[0] for l in labels if l.startswith('gpt-5')})
    order = ['gpt-4o-mini'] + [f'{f} ({e})' for f in fams for e in EFF if f'{f} ({e})' in tps]
    rows = []; thin = []
    for lab in order:
        n = len(tps[lab])
        if n < 8 and lab != 'gpt-4o-mini': thin.append((lab, n)); continue                          # F6
        dm = [st.median(v) for _, v in sorted(tps_day[lab].items())]
        c = R.get((lab, '1k', 'reason'), []); s = R.get((lab, 'short', 'reason'), [])
        l = R.get((lab, '10k', 'reason'), []); hv = R.get((lab, '1k', 'reason-heavy'), [])
        nonr = lab in NONREASON
        rows.append(dict(row=ROW_LABEL.get(lab, lab), label=lab, tps_p50=r2(med(tps[lab])), n_tps=n,
            tps_daily_mean=round(st.mean(dm), 1), tps_daily_sd=round(st.stdev(dm), 1) if len(dm) > 1 else None, days=len(dm),
            R_central=1.0 if nonr else r2(med(c)), n_R=len(c), R_short=1.0 if nonr else r2(med(s)), n_short=len(s),
            R_long=1.0 if nonr else r2(med(l)), n_long=len(l), R_floor=1.0 if nonr else r2(med(floor[lab])),
            R_heavy=1.0 if nonr else r2(med(hv)), n_heavy=len(hv), mode='None' if nonr else lab.split('(')[1][:-1].capitalize(),
            pqpc_class='4o' if lab == 'gpt-4o-mini' else '5x'))
    with open(f'{a.out}/corpus.csv', 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    with open(f'{a.out}/table_b1.csv', 'w', newline='') as f:
        w = csv.writer(f); w.writerow(['variant', 'tps_daily_mean', 'sd', 'days'])
        for r in rows: w.writerow([r['label'], r['tps_daily_mean'], r['tps_daily_sd'], r['days']])
    # ---- bootstrap bands (Appendix D.6) ----
    rng = np.random.default_rng(20260822)
    anchor_dm = np.array([st.median(v) for _, v in sorted(tps_day['gpt-4o-mini'].items())])
    boot = []
    for r in rows:
        lab = r['label']; dm = np.array([st.median(v) for _, v in sorted(tps_day[lab].items())])
        point = np.median(anchor_dm) / np.median(dm)
        ratios = np.median(anchor_dm[rng.integers(0, len(anchor_dm), (B, len(anchor_dm)))], axis=1) / \
                 np.median(dm[rng.integers(0, len(dm), (B, len(dm)))], axis=1)
        Tl, Th = np.percentile(ratios, 5) / point, np.percentile(ratios, 95) / point
        c = np.array(R.get((lab, '1k', 'reason'), []))
        if lab in NONREASON or len(c) == 0: Rl = Rh = 1.0; cm = 1.0; cn = 0
        else:
            cm = np.median(c); cn = len(c)
            bm = np.median(c[rng.integers(0, cn, (B, cn))], axis=1)
            Rl, Rh = np.percentile(bm, 5) / cm, np.percentile(bm, 95) / cm
        boot.append(dict(row=r['row'], label=lab, days=r['days'], T_low=round(Tl, 4), T_high=round(Th, 4),
                         R_n=cn, R_median=round(float(cm), 3), R_low=round(float(Rl), 4), R_high=round(float(Rh), 4)))
    with open(f'{a.out}/bootstrap.csv', 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(boot[0])); w.writeheader(); w.writerows(boot)
    # ---- ratio envelopes (Table C3) ----
    rng = np.random.default_rng(20260823); rat = []
    for num, den in RATIOS:
        def draw(lab):
            dm = np.array([st.median(v) for _, v in sorted(tps_day[lab].items())])
            c = np.array(R.get((lab, '1k', 'reason'), [])) if lab not in NONREASON else np.array([1.0])
            t = np.median(dm[rng.integers(0, len(dm), (B, len(dm)))], axis=1)
            rr = np.median(c[rng.integers(0, len(c), (B, len(c)))], axis=1) if len(c) > 1 else np.ones(B)
            return t, rr, np.median(dm), (np.median(c) if len(c) else 1.0)
        tn, rn_, tpn, rcn = draw(num); td, rd, tpd, rcd = draw(den)
        ratio = (rn_ / rd) * (td / tn)
        boot_point = (rcn / rcd) * (tpd / tpn)                       # bootstrap centre (median of daily medians)
        central = (rcn / rcd) * (med(tps[den]) / med(tps[num]))      # deposited central: Corpus P50 throughput
        lo, hi = np.percentile(ratio, 5) / boot_point, np.percentile(ratio, 95) / boot_point   # relative factors
        rat.append(dict(numerator=ROW_LABEL.get(num, num), denominator=ROW_LABEL.get(den, den), central=round(float(central), 3),
                        low=round(float(central * lo), 3), high=round(float(central * hi), 3),
                        low_factor=round(float(lo), 4), high_factor=round(float(hi), 4)))
    with open(f'{a.out}/ratios.csv', 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rat[0])); w.writeheader(); w.writerows(rat)
    json.dump(dict(exclusions=excl, thin_variants=thin, corpus_sha256=CORPUS_SHA, B=B, seeds=[20260822, 20260823]),
              open(f'{a.out}/provenance.json', 'w'), indent=1)
    print('variants:', len(rows), 'thin (excluded):', thin, 'exclusions:', excl)
    if a.workbook: compare(a.workbook, rows, boot, rat)

def compare(wb_path, rows, boot, rat):
    import openpyxl
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb['Corpus']; wbrows = {ws.cell(r, 2).value: [ws.cell(r, c).value for c in range(3, 17)] for r in range(3, 60) if ws.cell(r, 2).value}
    keys = ['tps_p50', 'n_tps', 'tps_daily_mean', 'tps_daily_sd', 'days', 'R_central', 'n_R', 'R_short', 'n_short', 'R_long', 'n_long', 'R_floor', 'R_heavy', 'n_heavy']
    bad = 0; tot = 0
    for r in rows:
        if r['label'] not in wbrows: print('not in workbook:', r['label']); continue
        for k, v in zip(keys, wbrows[r['label']]):
            tot += 1
            if v is None or r[k] is None: continue
            if abs(float(v) - float(r[k])) > 0.011 + (0 if k.startswith('n') or k == 'days' else 0):
                bad += 1; print(f'  DIFF {r["label"]:30s} {k:14s} workbook={v} corpus={r[k]}')
    print(f'Corpus sheet: {tot - bad}/{tot} cells agree')
    ws = wb['Bootstrap']; wbb = {ws.cell(r, 2).value: [ws.cell(r, c).value for c in (4, 5, 11, 12)] for r in range(3, 60) if ws.cell(r, 2).value}
    d = []
    for b in boot:
        if b['label'] in wbb:
            for k, v in zip(['T_low', 'T_high', 'R_low', 'R_high'], wbb[b['label']]):
                if v is not None: d.append(abs(float(v) - b[k]))
    print(f'Bootstrap sheet: {len(d)} bands compared, max |diff| = {max(d):.4f}, mean = {sum(d)/len(d):.4f}')
    ws = wb['RatioEnvelopes']; wr = {(ws.cell(r, 2).value, ws.cell(r, 3).value): [ws.cell(r, c).value for c in (10, 12, 13)] for r in range(3, 20) if ws.cell(r, 2).value}
    d = []
    for x in rat:
        v = wr.get((x['numerator'], x['denominator']))
        if not v: print('  ratio row not in workbook:', x['numerator'], x['denominator']); continue
        if v: d += [abs(float(v[0]) - x['central']), abs(float(v[1]) - x['low']), abs(float(v[2]) - x['high'])]
    print(f'RatioEnvelopes: {len(d)} values compared, max |diff| = {max(d):.3f}')

if __name__ == '__main__':
    main()
